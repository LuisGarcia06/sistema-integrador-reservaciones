from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Arquitectura_Tecnica_Implementada_2026-09-15_v2.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{margin}"))
        if element is None:
            element = OxmlElement(f"w:{margin}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_run(run, bold=False, italic=False, size=None, color="000000", font="Aptos"):
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)


def configure_document(doc, footer_text):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12), ("Heading 3", 10.5)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.font.size = Pt(size)
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(footer_text)
    style_run(run, size=8)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        style_run(run, bold=True)
    return paragraph


def add_para(doc, text="", lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.05
    if lead:
        lead_run = paragraph.add_run(lead)
        style_run(lead_run, bold=True)
        paragraph.add_run(" ")
    run = paragraph.add_run(text)
    style_run(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        style_run(run)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for index, text in enumerate(headers):
        cell = hdr.cells[index]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E5F")
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        style_run(run, bold=True, color="FFFFFF", size=9)
        if widths:
            cell.width = widths[index]

    for row_index, row_data in enumerate(rows):
        row = table.add_row()
        for col_index, text in enumerate(row_data):
            cell = row.cells[col_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, "F5F9FB")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(text))
            style_run(run, size=8.7)
            if widths:
                cell.width = widths[col_index]
    doc.add_paragraph()
    return table


def add_title(doc):
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Documentación de Arquitectura Técnica Implementada")
    style_run(run, bold=True, size=22)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Sistema Integrador de Reservaciones para Community Tours Sian Ka'an")
    style_run(run, size=12)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Revisión técnica: 15 de septiembre de 2026 | Versión documental: v2")
    style_run(run, size=9)


def build_document():
    doc = Document()
    configure_document(doc, "Sistema Integrador de Reservaciones - Arquitectura implementada v2")
    add_title(doc)

    add_heading(doc, "1. Introducción", 1)
    add_para(
        doc,
        "Este documento describe la arquitectura técnica realmente implementada al 15 de septiembre de 2026. La fuente de verdad es el código del repositorio: backend, frontend, Electron, PostgreSQL, package.json y migraciones SQL."
    )

    add_heading(doc, "2. Objetivo", 1)
    add_para(
        doc,
        "Servir como checkpoint técnico actualizado para continuar el cierre del sistema sin depender de documentos históricos que ya fueron superados por la implementación."
    )

    add_heading(doc, "3. Arquitectura General", 1)
    add_para(
        doc,
        "La aplicación conserva una arquitectura por capas: Electron contiene la experiencia de escritorio, el frontend vanilla consume la API REST, Express centraliza reglas y PostgreSQL persiste los datos. El frontend nunca accede directamente a PostgreSQL."
    )
    add_table(
        doc,
        ["Capa", "Implementación", "Responsabilidad"],
        [
            ["Electron", "electron/main.js y electron/preload.js", "Contenedor de escritorio, ciclo de vida de ventana y exportación PDF del Daily."],
            ["Frontend", "frontend/index.html, CSS y JavaScript modular", "Pantallas, navegación hash, consumo de API y renderizado operativo."],
            ["API", "Node.js + Express", "Rutas REST, autenticación, autorización, validación y respuestas HTTP."],
            ["Servicios", "backend/services", "Reglas de negocio, transacciones, capacidad y auditoría."],
            ["Datos", "PostgreSQL mediante backend/config/database.js", "Persistencia relacional con consultas parametrizadas."],
        ],
        [Inches(1.25), Inches(2.0), Inches(4.0)],
    )

    add_heading(doc, "4. Electron", 1)
    add_bullets(
        doc,
        [
            "El punto de entrada declarado en package.json es electron/main.js.",
            "Electron abre la aplicación local y expone funciones controladas mediante preload.",
            "La exportación PDF del Daily se realiza desde el proceso principal con el canal daily:export-pdf.",
            "El build Windows está configurado con electron-builder; el release final debe reconstruirse después de pruebas reales.",
        ],
    )

    add_heading(doc, "5. Frontend", 1)
    add_bullets(
        doc,
        [
            "El frontend usa HTML, CSS y JavaScript sin framework.",
            "La navegación usa rutas hash como #/dashboard, #/reservaciones, #/operaciones, #/daily, #/historial y #/bitacora.",
            "Todas las operaciones de datos usan fetch contra /api.",
            "Las pantallas implementadas cubren Login, Dashboard, Reservaciones, Operaciones, Daily, Historial, Bitácora, Usuarios y Configuración.",
        ],
    )

    add_heading(doc, "6. API REST", 1)
    add_table(
        doc,
        ["Ruta", "Uso principal", "Acceso"],
        [
            ["/api/auth", "Login", "Público para login"],
            ["/api/dashboard", "Resumen operativo real", "Usuarios autenticados"],
            ["/api/reservaciones", "Listado, creación, edición, cancelación y asignación a transporte", "Consulta lee; Administrador modifica"],
            ["/api/operaciones", "Operaciones, detalle y sugerencias por fecha", "Consulta lee; Administrador modifica"],
            ["/api/transportes", "Transportes por operación", "Consulta lee; Administrador modifica"],
            ["/api/daily", "Daily operativo y observaciones por fecha", "Usuarios autenticados"],
            ["/api/bitacora", "Auditoría consultable", "Administrador"],
            ["/api/usuarios", "Usuarios y roles", "Administrador"],
            ["/api/tours, /api/paises, /api/plataformas", "Catálogos base", "Según ruta y rol"],
            ["/api/guias, /api/operadores, /api/vehiculos", "Catálogos operativos", "Según ruta y rol"],
        ],
        [Inches(2.0), Inches(3.2), Inches(2.0)],
    )

    add_heading(doc, "7. PostgreSQL", 1)
    add_table(
        doc,
        ["Tabla", "Propósito"],
        [
            ["roles", "Catálogo de permisos funcionales."],
            ["usuarios", "Cuentas de acceso con rol, correo, contraseña hasheada y estado."],
            ["tours", "Catálogo de tours vendibles."],
            ["paises", "Catálogo de nacionalidades o países del cliente."],
            ["plataformas", "Origen comercial de la reservación."],
            ["vehiculos", "Unidades disponibles, placas, color, capacidad y estado."],
            ["operadores", "Conductores u operadores de transporte."],
            ["guias", "Guías asignables a grupos operativos."],
            ["operaciones_tour", "Grupo operativo por fecha, tour, turno y número de grupo."],
            ["transportes_operacion", "Transporte asociado a un grupo operativo."],
            ["reservaciones", "Datos comerciales, pickup, turno, estado y asignación operativa."],
            ["daily_observaciones", "Observaciones generales por fecha del Daily."],
            ["bitacora", "Registro de auditoría asociado a usuario y reservación."],
        ],
        [Inches(2.0), Inches(5.2)],
    )

    add_heading(doc, "8. Autenticación y Autorización", 1)
    add_bullets(
        doc,
        [
            "El login se realiza mediante POST /api/auth/login.",
            "Las contraseñas se validan con bcryptjs.",
            "La sesión de API usa JWT.",
            "La autorización por rol se valida en middleware backend.",
            "El rol Administrador modifica información; el rol Consulta conserva acceso de lectura según módulo.",
        ],
    )

    add_heading(doc, "9. Reservaciones", 1)
    add_bullets(
        doc,
        [
            "Las nuevas reservaciones se registran como Pendiente.",
            "Los estados documentados son Pendiente, Confirmada, Activa, Cancelada y Completada.",
            "El turno es explícito: Mañana o Tarde.",
            "Puede existir turno NULL solo para datos heredados.",
            "No se documentan transiciones automáticas de estado.",
            "pickup_time pertenece a cada reservación.",
            "La asignación a transporte es manual y valida fecha, tour, turno y capacidad.",
        ],
    )

    add_heading(doc, "10. Operaciones y Grupos", 1)
    add_table(
        doc,
        ["Relación o restricción", "Regla vigente"],
        [
            ["roles 1:N usuarios", "Cada usuario pertenece a un rol."],
            ["tours 1:N reservaciones", "Cada reservación corresponde a un tour."],
            ["tours 1:N operaciones_tour", "Cada grupo operativo corresponde a un tour."],
            ["guias 1:N operaciones_tour", "Un guía puede aparecer en varios grupos."],
            ["operaciones_tour 0:1 transportes_operacion", "Cada grupo puede tener como máximo un transporte."],
            ["vehiculos 1:N transportes_operacion", "Un vehículo puede usarse en distintos grupos."],
            ["operadores 1:N transportes_operacion", "Un operador puede usarse en distintos grupos."],
            ["transportes_operacion 1:N reservaciones", "Varias reservaciones pueden asignarse al transporte del grupo."],
            ["reservaciones 1:N bitacora", "Cada cambio relevante de reservación queda trazado."],
            ["fecha + tour + turno + grupo", "Identidad única de operación."],
        ],
        [Inches(2.5), Inches(4.7)],
    )
    add_bullets(
        doc,
        [
            "Una operación representa un grupo.",
            "La identidad lógica de una operación es fecha + tour + turno + numero_grupo.",
            "Los grupos permitidos son 1 y 2.",
            "Cada grupo admite hasta 12 PAX activos.",
            "Un tour en una fecha y turno admite hasta 24 PAX activos.",
            "hora_inicio representa la primera hora de pickup del grupo.",
        ],
    )

    add_heading(doc, "11. Transporte, Guía y Operador", 1)
    add_bullets(
        doc,
        [
            "Cada grupo puede tener un guía asignado.",
            "Cada grupo puede tener como máximo un transporte.",
            "El transporte puede asociar vehículo, operador y observaciones del operador.",
            "Un transporte con vehículo debe respetar reglas de capacidad y mínimo operativo cuando ya tiene reservas activas.",
        ],
    )

    add_heading(doc, "12. Operaciones Sugeridas", 1)
    add_table(
        doc,
        ["Caso", "Resultado"],
        [
            ["Agrupación", "Fecha + tour + turno."],
            ["Estados incluidos", "Pendiente, Confirmada y Activa."],
            ["Estados excluidos", "Cancelada y Completada."],
            ["1 a 12 PAX", "Sugiere 1 grupo."],
            ["13 a 24 PAX", "Sugiere 2 grupos."],
            [">24 PAX", "Mantiene máximo 2 grupos y reporta excedente."],
            ["Asignación", "No distribuye clientes automáticamente; el usuario asigna según hotel, ruta y pickup."],
        ],
        [Inches(2.0), Inches(5.2)],
    )

    add_heading(doc, "13. Daily", 1)
    add_para(
        doc,
        "El Daily operativo muestra salidas por fecha con Tour, Turno, Grupo, Fecha, Primera hora de pickup, Guía, Vehículo, Operador y PAX. Las reservaciones se ordenan por pickup_time."
    )
    add_bullets(
        doc,
        [
            "GET /api/daily/operativo entrega operaciones, transportes y reservaciones sin asignar.",
            "GET/PUT /api/daily/observaciones administra observaciones por fecha.",
            "La impresión y PDF generan hoja por grupo/transporte cuando corresponde.",
        ],
    )

    add_heading(doc, "14. Dashboard", 1)
    add_table(
        doc,
        ["Métrica", "Fuente"],
        [
            ["Reservaciones vigentes de hoy", "GET /api/dashboard/resumen"],
            ["PAX vigentes", "GET /api/dashboard/resumen"],
            ["Grupos preparados", "GET /api/dashboard/resumen"],
            ["Grupos pendientes", "GET /api/dashboard/resumen"],
            ["Reservaciones sin turno", "GET /api/dashboard/resumen"],
            ["Alertas de capacidad", "GET /api/dashboard/resumen"],
            ["Últimas 10 reservaciones", "GET /api/dashboard/resumen"],
            ["Salidas de hoy", "GET /api/dashboard/resumen"],
        ],
        [Inches(3.0), Inches(4.2)],
    )

    add_heading(doc, "15. Historial y Bitácora", 1)
    add_bullets(
        doc,
        [
            "Historial consulta reservaciones con filtros desde /api/reservaciones.",
            "Bitácora se consulta desde /api/bitacora y está protegida para Administrador.",
            "Crear, modificar, cancelar y asignar transportes registra auditoría cuando aplica.",
        ],
    )

    add_heading(doc, "16. Seguridad", 1)
    add_bullets(
        doc,
        [
            "No se exponen credenciales PostgreSQL al frontend.",
            "Las consultas SQL usan parámetros.",
            "Las contraseñas se almacenan hasheadas.",
            "Los permisos se validan en backend.",
            "Los secretos deben vivir en .env.",
        ],
    )

    add_heading(doc, "17. Configuración Runtime", 1)
    add_bullets(
        doc,
        [
            "La configuración sensible se toma de variables de entorno.",
            "La conexión PostgreSQL se centraliza en backend/config/database.js.",
            "El backend Express sirve el frontend en /app.",
            "package.json define start, desktop, pack:win y dist:win.",
        ],
    )

    add_heading(doc, "18. Installer", 1)
    add_para(
        doc,
        "Existe configuración técnica de empaquetado con electron-builder. El release final queda pendiente de reconstrucción después del cierre documental, datos reales y pruebas de aceptación."
    )

    add_heading(doc, "19. Integraciones Pendientes", 1)
    add_para(
        doc,
        "FareHarbor, GetYourGuide y WhatsApp Business están contempladas como plataformas u origen de reservación, pero no existe sincronización automática implementada ni contratos técnicos autorizados en el código."
    )

    add_heading(doc, "20. Estado Técnico Actual", 1)
    add_table(
        doc,
        ["Área", "Estado", "Observación"],
        [
            ["Backend", "Implementado", "API REST, servicios, validadores, seguridad y auditoría."],
            ["Frontend", "Implementado", "Pantallas operativas en JavaScript vanilla."],
            ["Base de datos", "Implementada", "Schema actual más migraciones 001 a 007."],
            ["Daily y Dashboard", "Implementados", "Ambos consumen datos reales del backend."],
            ["Integraciones", "Pendiente", "Requieren documentación y acceso externo."],
            ["Pruebas reales", "Pendiente", "Faltan datos reales, aceptación de usuarios e installer final."],
        ],
        [Inches(1.7), Inches(1.4), Inches(4.1)],
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Anexo A. Comandos de Desarrollo", 1)
    add_table(
        doc,
        ["Comando", "Descripción"],
        [
            ["npm start", "Inicia el backend Express y sirve el frontend en /app."],
            ["npm run desktop", "Inicia Electron sobre el backend local."],
            ["npm run pack:win", "Genera paquete Windows sin instalador final."],
            ["npm run dist:win", "Genera instalador Windows con electron-builder."],
            ["python scripts/generar_documento_arquitectura_implementada.py", "Genera este documento en versión nueva."],
        ],
        [Inches(2.7), Inches(4.5)],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(f"Documento generado: {path}")
